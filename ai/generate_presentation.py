"""
Enterprise Deck Generator — Foundation Platform
Premium dark consulting-grade presentation based on storyboard.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ─────────────────────────────────────────────
# PALETTE (Premium Dark Theme)
# ─────────────────────────────────────────────
BG_DEEP_NAVY = RGBColor(0x0B, 0x10, 0x20)
SURFACE = RGBColor(0x11, 0x18, 0x27)
SURFACE_SEC = RGBColor(0x1E, 0x29, 0x3B)

TEXT_PRIMARY = RGBColor(0xF8, 0xFA, 0xFC)
TEXT_SECONDARY = RGBColor(0xCB, 0xD5, 0xE1)
TEXT_MUTED = RGBColor(0x94, 0xA3, 0xB8)

ACCENT_BLUE = RGBColor(0x38, 0xBD, 0xF8)
ACCENT_GREEN = RGBColor(0x22, 0xC5, 0x5E)
ACCENT_VIOLET = RGBColor(0x8B, 0x5C, 0xF6)
ACCENT_AMBER = RGBColor(0xF5, 0x9E, 0x0B)
ACCENT_ERROR = RGBColor(0xEF, 0x44, 0x44)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)
HEADER_H = Inches(1.5)

# ─────────────────────────────────────────────
# CORE HELPERS
# ─────────────────────────────────────────────
def new_prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs

def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color

def rect(slide, l, t, w, h, fill, line=None, lw=1.0):
    sh = slide.shapes.add_shape(1, l, t, w, h)
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh

def txt(slide, text, l, t, w, h, size=14, bold=False, color=TEXT_PRIMARY, align=PP_ALIGN.LEFT, font="Arial", italic=False):
    b = slide.shapes.add_textbox(l, t, w, h)
    b.word_wrap = True
    tf = b.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return b

def multiline(slide, lines, l, t, w, h, size=12, color=TEXT_SECONDARY, font="Arial", line_gap=None):
    b = slide.shapes.add_textbox(l, t, w, h)
    b.word_wrap = True
    tf = b.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if line_gap and i > 0:
            p.space_before = Pt(line_gap)
        r = p.add_run()
        r.text = line
        r.font.name = font
        r.font.size = Pt(size)
        r.font.color.rgb = color
    return b

def add_notes(slide, note_text):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = note_text

def header(slide, title, align=PP_ALIGN.LEFT):
    txt(slide, title, Inches(0.8), Inches(0.4), SLIDE_W - Inches(1.6), Inches(0.8), size=36, bold=True, color=TEXT_PRIMARY, align=align)

def footer(slide, slide_num):
    txt(slide, "foundation-platform", Inches(0.8), SLIDE_H - Inches(0.6), Inches(3), Inches(0.4), size=10, color=TEXT_MUTED)
    txt(slide, str(slide_num), SLIDE_W - Inches(1.2), SLIDE_H - Inches(0.6), Inches(0.5), Inches(0.4), size=10, color=TEXT_MUTED, align=PP_ALIGN.RIGHT)

def card(slide, top, left, width, height, title, content_lines, accent_color=ACCENT_BLUE):
    rect(slide, left, top, width, height, SURFACE, line=SURFACE_SEC)
    rect(slide, left, top, width, Inches(0.06), accent_color)
    txt(slide, title, left + Inches(0.3), top + Inches(0.2), width - Inches(0.6), Inches(0.5), size=18, bold=True, color=TEXT_PRIMARY)
    multiline(slide, content_lines, left + Inches(0.3), top + Inches(0.8), width - Inches(0.6), height - Inches(1.0), size=14, color=TEXT_SECONDARY, line_gap=8)

# ─────────────────────────────────────────────
# SLIDE TEMPLATES
# ─────────────────────────────────────────────

def slide_1_title(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    rect(slide, Inches(2), Inches(3), Inches(1.5), Inches(0.1), ACCENT_BLUE)
    txt(slide, "foundation-platform", Inches(2), Inches(3.2), Inches(9), Inches(1.2), size=56, bold=True)
    txt(slide, "A faster way to build production-ready Spring Boot services", Inches(2), Inches(4.5), Inches(9), Inches(0.6), size=24, color=TEXT_SECONDARY)
    txt(slide, "Fast. Standardized. Production-ready.", Inches(2), Inches(5.5), Inches(9), Inches(0.5), size=16, color=ACCENT_GREEN, bold=True)
    add_notes(slide, "Aujourd'hui, l'objectif n'est pas juste de présenter un ensemble de modules Maven ou de starters Spring Boot. L'objectif est de montrer comment nous pouvons industrialiser la création de microservices Spring Boot, réduire la duplication, et fournir une base prête pour la production cohérente pour les futurs services.\n\nCe socle est un accélérateur stratégique.")
    return slide

def slide_2_challenge(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Every new service starts with the same hidden work")
    txt(slide, "Before business logic, teams rebuild the same technical foundation.", Inches(0.8), Inches(1.2), Inches(10), Inches(0.5), size=18, color=TEXT_SECONDARY)
    
    blocks = ["Security", "Logging", "Observability", "API errors", "Data setup", "Messaging", "Clients", "Tests"]
    y_start = Inches(2.5)
    x_start = Inches(2)
    for i, blk in enumerate(blocks):
        row = i // 4
        col = i % 4
        x = x_start + col * Inches(2.5)
        y = y_start + row * Inches(1.5)
        rect(slide, x, y, Inches(2), Inches(1), SURFACE, line=SURFACE_SEC)
        txt(slide, blk, x, y + Inches(0.3), Inches(2), Inches(0.4), align=PP_ALIGN.CENTER, size=14)
        
    corner_txt = "At the center, the developer struggles to start..."
    txt(slide, corner_txt, Inches(0.8), Inches(6), Inches(11), Inches(0.5), size=16, color=TEXT_MUTED, align=PP_ALIGN.CENTER, italic=True)

    add_notes(slide, "Lorsqu'une équipe démarre un nouveau microservice, elle commence rarement tout de suite par la logique métier. En premier lieu, il faut la sécurité, les logs, l'ID de corrélation, la gestion des erreurs API, l'observabilité, la configuration des données, le messaging, les clients et les tests. Ces sujets sont obligatoires, mais ils sont répétitifs et souvent réimplémentés différemment par chaque équipe.")
    footer(slide, 2)
    return slide

def slide_3_cost(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "The real cost is not only time. It is inconsistency.")
    
    lines = [
        "• Duplicated setup",
        "• Different patterns",
        "• Harder reviews",
        "• Uneven production readiness",
        "• Slower onboarding"
    ]
    multiline(slide, lines, Inches(1), Inches(2.2), Inches(5), Inches(4), size=20, line_gap=16)
    
    rect(slide, Inches(6.5), Inches(2.2), Inches(2.8), Inches(3.5), SURFACE_SEC, line=ACCENT_ERROR)
    rect(slide, Inches(9.5), Inches(2.2), Inches(2.8), Inches(3.5), SURFACE_SEC, line=ACCENT_AMBER)
    txt(slide, "Service A\n(Custom Security, weak logs)", Inches(6.5), Inches(3.5), Inches(2.8), Inches(1), align=PP_ALIGN.CENTER, color=TEXT_SECONDARY)
    txt(slide, "Service B\n(Different Error formats)", Inches(9.5), Inches(3.5), Inches(2.8), Inches(1), align=PP_ALIGN.CENTER, color=TEXT_SECONDARY)
    
    add_notes(slide, "Le coût n'est pas seulement le temps passé au début de chaque service. Le coût profond, c'est l'incohérence. Si chaque équipe implémente ses propres erreurs API, son mapping de sécurité, sa stratégie de logs, l'observabilité et les conventions de données, nous créons une friction à long terme pour la gouvernance de l'architecture, le support, les audits et les opérations.")
    footer(slide, 3)
    return slide

def slide_4_foundation(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "From repeated setup to a governed foundation")
    
    lines = [
        "One foundation.",
        "Reusable capabilities.",
        "Consistent services.",
        "Faster delivery."
    ]
    multiline(slide, lines, Inches(1), Inches(3), Inches(4), Inches(3), size=20, color=ACCENT_BLUE, line_gap=16)
    
    rect(slide, Inches(5.5), Inches(3.2), Inches(2.5), Inches(1), SURFACE)
    txt(slide, "Repeated setup", Inches(5.5), Inches(3.5), Inches(2.5), Inches(0.4), align=PP_ALIGN.CENTER)
    
    txt(slide, "→", Inches(8), Inches(3.5), Inches(0.5), Inches(0.4), size=24, color=ACCENT_BLUE)
    
    rect(slide, Inches(8.5), Inches(2.8), Inches(2.5), Inches(1.8), SURFACE, line=ACCENT_BLUE, lw=2)
    txt(slide, "foundation-platform", Inches(8.5), Inches(3.5), Inches(2.5), Inches(0.4), align=PP_ALIGN.CENTER, color=ACCENT_BLUE, bold=True)
    
    txt(slide, "→", Inches(11), Inches(3.5), Inches(0.5), Inches(0.4), size=24, color=ACCENT_BLUE)
    
    rect(slide, Inches(11.5), Inches(2.5), Inches(1.5), Inches(0.6), SURFACE)
    rect(slide, Inches(11.5), Inches(3.4), Inches(1.5), Inches(0.6), SURFACE)
    rect(slide, Inches(11.5), Inches(4.3), Inches(1.5), Inches(0.6), SURFACE)
    txt(slide, "Standardized\nServices", Inches(11.5), Inches(2.8), Inches(1.5), Inches(1.5), align=PP_ALIGN.CENTER, size=12)

    add_notes(slide, "foundation-platform transforme les sujets techniques récurrents en capacités réutilisables. L'approche n'est pas de créer un framework propriétaire lourd. L'approche est de rester proche de Spring Boot tout en packageant les conventions communes dans une fondation propre et réutilisable.")
    footer(slide, 4)
    return slide

def slide_5_principles(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Designed to accelerate, not to lock in")
    
    cards = [
        ("Lightweight", "No heavy abstractions"),
        ("Spring Boot-native", "Standard ecosystem"),
        ("Dependency-driven", "Pick only what you need"),
        ("Override-friendly", "Customizable defaults"),
        ("IDP-agnostic", "Any OIDC provider"),
        ("No business logic", "Purely technical")
    ]
    
    for i, (title, sub) in enumerate(cards):
        col = i % 3
        row = i // 3
        card(slide, Inches(2.5) + row * Inches(2.2), Inches(1) + col * Inches(3.8), Inches(3.4), Inches(1.8), title, [sub], ACCENT_BLUE if i%2==0 else ACCENT_GREEN)

    add_notes(slide, "Les principes de conception sont essentiels pour l'adoption. foundation-platform est léger, natif Spring Boot, piloté par les dépendances et facile à surcharger (override). Il ne contient aucune logique métier. Il n'est pas couplé à un Fournisseur d'Identité spécifique. Il est conçu pour standardiser sans enfermer les équipes dans un modèle de programmation propriétaire.")
    footer(slide, 5)
    return slide

def slide_6_architecture(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "A modular foundation")
    
    rect(slide, Inches(2), Inches(2), Inches(9.33), Inches(0.8), SURFACE, line=ACCENT_BLUE)
    txt(slide, "Layer 1 : Service generator\nfoundation-archetype", Inches(2), Inches(2.2), Inches(9.33), Inches(0.6), align=PP_ALIGN.CENTER, size=14)

    rect(slide, Inches(2), Inches(3.1), Inches(9.33), Inches(0.8), SURFACE, line=ACCENT_GREEN)
    txt(slide, "Layer 2 : Mandatory capabilities\ncore | api | security | logging | observability | mapping | test", Inches(2), Inches(3.3), Inches(9.33), Inches(0.6), align=PP_ALIGN.CENTER, size=14)
    
    rect(slide, Inches(2), Inches(4.2), Inches(9.33), Inches(0.8), SURFACE, line=ACCENT_AMBER)
    txt(slide, "Layer 3 : Optional capabilities\ndata | nats | http-client | soap-client", Inches(2), Inches(4.4), Inches(9.33), Inches(0.6), align=PP_ALIGN.CENTER, size=14)
    
    rect(slide, Inches(2), Inches(5.3), Inches(9.33), Inches(0.8), SURFACE, line=TEXT_MUTED)
    txt(slide, "Layer 4 : Build baseline\nparent | bom | common", Inches(2), Inches(5.5), Inches(9.33), Inches(0.6), align=PP_ALIGN.CENTER, size=14)

    add_notes(slide, "La plateforme est structurée autour de quatre groupes. En bas, nous avons la base de build avec le parent, le BOM et les composants communs. Au-dessus, les capacités obligatoires que chaque service reçoit. Ensuite, les capacités optionnelles ne sont sélectionnées qu'en cas de besoin. Enfin, l'archetype fournit le point d'entrée développeur pour générer un service.")
    footer(slide, 6)
    return slide

def slide_7_dx(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "A new service in minutes")
    txt(slide, "Generate. Select capabilities. Start coding.", Inches(0.8), Inches(1.2), Inches(10), Inches(0.5), size=18, color=ACCENT_GREEN)
    
    # Left Terminal
    rect(slide, Inches(0.8), Inches(2.2), Inches(6), Inches(3.5), SURFACE, line=SURFACE_SEC)
    rect(slide, Inches(0.8), Inches(2.2), Inches(6), Inches(0.4), SURFACE_SEC)
    txt(slide, ">_", Inches(1), Inches(2.25), Inches(1), Inches(0.3), size=12, color=TEXT_SECONDARY)
    
    cmd = "mvn archetype:generate \\\n  -DartifactId=my-service \\\n  -DserviceName=MyService \\\n  -Dcapabilities=data,nats"
    multiline(slide, [cmd], Inches(1), Inches(2.8), Inches(5.5), Inches(2.5), size=13, font="Consolas", color=ACCENT_BLUE)
    
    # Right Blueprint
    rect(slide, Inches(7.5), Inches(2.2), Inches(5), Inches(3.5), SURFACE, line=ACCENT_BLUE)
    lines = [
        "✓ Hexagonal structure",
        "✓ Mandatory starters",
        "✓ Selected capabilities (data, nats)",
        "✓ Minimal configuration",
        "✓ Ready for business code"
    ]
    multiline(slide, lines, Inches(7.8), Inches(2.5), Inches(4.4), Inches(3), size=16, line_gap=12)

    add_notes(slide, "L'expérience développeur est simple. Une équipe génère un nouveau service, sélectionne les capacités nécessaires, et obtient un projet prêt à coder avec la bonne architecture hexagonale et la base de référence. L'équipe peut se concentrer sur la logique métier au lieu de réinventer la configuration technique.")
    footer(slide, 7)
    return slide

def slide_8_mandatory(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Every service starts with the right baseline")
    
    caps = ["Core", "API", "Security", "Logging", "Observability", "Mapping", "Testing"]
    for i, cap in enumerate(caps):
        rect(slide, Inches(0.8) + i * Inches(1.7), Inches(3), Inches(1.5), Inches(2), SURFACE, line=ACCENT_BLUE)
        txt(slide, cap, Inches(0.8) + i * Inches(1.7), Inches(4), Inches(1.5), Inches(0.5), align=PP_ALIGN.CENTER, size=16)
        
    add_notes(slide, "Chaque service généré inclut automatiquement la fondation obligatoire : core, conventions API, sécurité, logging, observabilité, mapping et tests. Cela signifie que même le service le plus simple démarre avec des standards d'entreprise cohérents et prêts pour la production.")
    footer(slide, 8)
    return slide

def slide_9_optional(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Add only what the service needs")
    
    card(slide, Inches(2), Inches(1), Inches(5), Inches(2), "data", ["JPA + Flyway + PostgreSQL"], ACCENT_GREEN)
    card(slide, Inches(2), Inches(7), Inches(5), Inches(2), "nats", ["Messaging conventions"], ACCENT_VIOLET)
    card(slide, Inches(4.5), Inches(1), Inches(5), Inches(2), "http-client", ["WebClient + OpenAPI clients"], ACCENT_BLUE)
    card(slide, Inches(4.5), Inches(7), Inches(5), Inches(2), "soap-client", ["Apache CXF + WSDL clients"], ACCENT_AMBER)
    
    add_notes(slide, "Les capacités optionnelles ne sont sélectionnées que lorsque le service en a besoin. Un service d'API stateless n'a pas besoin de transporter de dépendances de données. Un service sans messagerie n'a pas besoin de NATS. Un service qui ne consomme pas de SOAP n'a pas besoin de CXF. Cela maintient la fondation modulaire et alignée sur les besoins réels du service, évitant l'effet 'gros framework'.")
    footer(slide, 9)
    return slide

def slide_10_migration(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Built for the migration journey")
    
    # Flow
    rect(slide, Inches(3), Inches(2), Inches(7.33), Inches(0.8), SURFACE)
    txt(slide, "webMethods / legacy integration", Inches(3), Inches(2.2), Inches(7.33), Inches(0.6), align=PP_ALIGN.CENTER, size=16, color=TEXT_MUTED)
    
    txt(slide, "↓", Inches(6.5), Inches(2.9), Inches(0.5), Inches(0.5), size=24, align=PP_ALIGN.CENTER, color=ACCENT_BLUE)
    
    rect(slide, Inches(3), Inches(3.5), Inches(7.33), Inches(1), BG_DEEP_NAVY, line=ACCENT_BLUE, lw=2)
    txt(slide, "foundation-platform", Inches(3), Inches(3.8), Inches(7.33), Inches(0.6), align=PP_ALIGN.CENTER, size=20, bold=True, color=ACCENT_BLUE)
    
    txt(slide, "↓", Inches(6.5), Inches(4.6), Inches(0.5), Inches(0.5), size=24, align=PP_ALIGN.CENTER, color=ACCENT_GREEN)
    
    rect(slide, Inches(3), Inches(5.2), Inches(7.33), Inches(0.8), SURFACE)
    txt(slide, "Spring Boot services", Inches(3), Inches(5.4), Inches(7.33), Inches(0.6), align=PP_ALIGN.CENTER, size=16)
    
    rect(slide, Inches(1), Inches(6.5), Inches(11.33), Inches(0.5), SURFACE_SEC)
    txt(slide, "REST APIs  |  SOAP clients  |  NATS  |  Flyway  |  Gravitee-ready", Inches(1), Inches(6.6), Inches(11.33), Inches(0.4), align=PP_ALIGN.CENTER, size=14, color=ACCENT_BLUE)

    add_notes(slide, "La fondation est conçue pour un parcours de migration. Elle supporte les API REST modernes, les intégrations SOAP legacy via des clients WSDL, la messagerie NATS, les migrations de schémas Flyway, et les services s'exécutant derrière Gravitee. Cela permet une migration progressive plutôt qu'un big bang risqué.")
    footer(slide, 10)
    return slide

def slide_11_governance(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Governance built into developer experience")
    
    rect(slide, Inches(3), Inches(2), Inches(7.33), Inches(4.5), SURFACE, line=ACCENT_BLUE)
    lines = [
        "☑ JWT / OAuth2 validation",
        "☑ Correlation ID propagation",
        "☑ Standard API errors",
        "☑ Flyway schema versions",
        "☑ Actuator health & metrics",
        "☑ OpenTelemetry tracing",
        "☑ MapStruct conventions",
        "☑ Standardized tests structure"
    ]
    multiline(slide, lines, Inches(4), Inches(2.5), Inches(5.33), Inches(3.5), size=18, line_gap=8)
    
    add_notes(slide, "La gouvernance ne doit pas vivre uniquement dans des documents Word. Elle doit être intégrée dans l'expérience du développeur. Avec foundation-platform, de nombreuses normes d'architecture et de production sont incluses par défaut dans le service généré et les starters. La gouvernance devient automatique.")
    footer(slide, 11)
    return slide

def slide_12_productivity(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "The productivity gain is repeatability")
    
    # 3 columns
    card(slide, Inches(2), Inches(1), Inches(3.5), Inches(3.5), "Before", ["Manual setup", "Duplicated code", "Uneven quality"], TEXT_MUTED)
    txt(slide, "→", Inches(4.7), Inches(3), Inches(0.5), Inches(0.5), size=24, color=TEXT_MUTED)
    
    card(slide, Inches(2), Inches(4.9), Inches(3.5), Inches(3.5), "With foundation", ["Generated baseline", "Shared standards", "Selected capabilities"], ACCENT_BLUE)
    txt(slide, "→", Inches(8.5), Inches(3), Inches(0.5), Inches(0.5), size=24, color=ACCENT_BLUE)
    
    card(slide, Inches(2), Inches(8.8), Inches(3.5), Inches(3.5), "Impact", ["Faster start", "Less rework", "Lower risk"], ACCENT_GREEN)
    
    add_notes(slide, "Le gain de productivité n'est pas seulement dû à la vitesse initiale. Le vrai gain est la répétabilité. Chaque nouveau service démarre avec la même base de référence, les mêmes conventions et le même niveau de qualité. Cela réduit les retouches, accélère les revues de code et diminue les risques en production.")
    footer(slide, 12)
    return slide

def slide_13_boundaries(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Clear boundaries make adoption easier")
    
    card(slide, Inches(2.5), Inches(1.5), Inches(4.5), Inches(3.5), "It is", ["Feature composition", "Spring Boot foundation", "Maven archetype", "Reusable starters", "Standardization layer", "Delivery accelerator"], ACCENT_GREEN)
    card(slide, Inches(2.5), Inches(7.33), Inches(4.5), Inches(3.5), "It is NOT", ["Heavy framework", "Business framework", "Deployment platform", "Specific IDP", "Replacement for Spring Boot"], ACCENT_ERROR)

    add_notes(slide, "L'adoption devient plus facile lorsque les limites sont claires. foundation-platform ne remplace pas Spring Boot, ne devient pas la plateforme de déploiement, n'intègre pas de logique métier et ne se couple pas à un fournisseur d'identité spécifique. Il se concentre exclusivement sur la base technique des services.")
    footer(slide, 13)
    return slide

def slide_14_adoption(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Adoption can be progressive")
    
    steps = [
        "1. Generate service",
        "2. Select capabilities",
        "3. Implement business logic",
        "4. Run standardized tests",
        "5. Deliver through existing CI/CD"
    ]
    
    for i, step in enumerate(steps):
        x = Inches(0.5) + i * Inches(2.5)
        rect(slide, x, Inches(3.5), Inches(2.2), Inches(1), SURFACE, line=ACCENT_BLUE)
        txt(slide, step, x + Inches(0.1), Inches(3.7), Inches(2), Inches(0.6), size=12, align=PP_ALIGN.CENTER)
        if i < 4:
            txt(slide, "→", x + Inches(2.2), Inches(3.8), Inches(0.3), Inches(0.4), size=18, color=ACCENT_BLUE)
            
    add_notes(slide, "Le modèle d'adoption est progressif. Les équipes peuvent commencer par un service généré, sélectionner uniquement les capacités requises, implémenter la logique métier et utiliser la chaîne de livraison existante (CI/CD). La fondation soutient l'adoption sans imposer une transformation 'big bang'.")
    footer(slide, 14)
    return slide

def slide_15_closing(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    
    rect(slide, Inches(1), Inches(3), Inches(0.1), Inches(2), ACCENT_BLUE)
    txt(slide, "foundation-platform is the accelerator\nfor the next wave of services", Inches(1.5), Inches(2.8), Inches(10), Inches(1.5), size=32, bold=True)
    
    lines = [
        "Build faster.",
        "Stay consistent.",
        "Keep Spring Boot native."
    ]
    multiline(slide, lines, Inches(1.5), Inches(4.5), Inches(5), Inches(2), size=24, color=ACCENT_GREEN, line_gap=12)

    add_notes(slide, "foundation-platform réduit le coût de démarrage, le risque de divergence et l'effort de gouvernance des services Spring Boot. Construisez plus vite. Restez cohérents. Gardez la philosophie native Spring Boot. La prochaine étape est de valider l'adoption à travers un service pilote, de recueillir les retours et de passer à l'échelle progressivement.")
    footer(slide, 15)
    return slide

# ─────────────────────────────────────────────
# BACKUP SLIDES
# ─────────────────────────────────────────────
def slide_b1_map(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Backup: Module map")
    
    txt(slide, "Build Foundation: parent, bom, common", Inches(1), Inches(2), Inches(8), Inches(0.5), size=14, color=TEXT_SECONDARY)
    txt(slide, "Mandatory: core, api, security, logging, observability, mapping, test", Inches(1), Inches(3), Inches(10), Inches(0.5), size=14, color=TEXT_SECONDARY)
    txt(slide, "Optional: data, nats, http-client, soap-client", Inches(1), Inches(4), Inches(8), Inches(0.5), size=14, color=TEXT_SECONDARY)
    txt(slide, "Generator: archetype", Inches(1), Inches(5), Inches(8), Inches(0.5), size=14, color=TEXT_SECONDARY)
    footer(slide, "B1")
    return slide

def slide_b2_archetype(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Backup: Service generation model")
    
    cmd = "mvn archetype:generate \\\n  -DarchetypeGroupId=fr.francetv.foundation \\\n  -DarchetypeArtifactId=foundation-archetype \\\n  -DartifactId=my-service \\\n  -DserviceName=MyService \\\n  -Dcapabilities=data,nats"
    rect(slide, Inches(1), Inches(2), Inches(5.5), Inches(2.5), SURFACE)
    multiline(slide, [cmd], Inches(1.2), Inches(2.2), Inches(5), Inches(2.1), size=11, font="Consolas", color=ACCENT_BLUE)
    
    rect(slide, Inches(7), Inches(2), Inches(5), Inches(2.5), SURFACE, line=ACCENT_GREEN)
    lines = ["Generated: mandatory + optional capabilities", "Hexagonal structure", "Minimal application.yml", "Optional Dockerfile / GitLab CI templates"]
    multiline(slide, lines, Inches(7.2), Inches(2.2), Inches(4.6), Inches(2.1), size=12)
    footer(slide, "B2")
    return slide

def slide_b3_hexagonal(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Backup: Generated hexagonal architecture")
    
    rect(slide, Inches(5), Inches(2.5), Inches(3.3), Inches(2.5), SURFACE, line=ACCENT_GREEN)
    txt(slide, "Domain\n(model, port in/out, service)", Inches(5), Inches(3.3), Inches(3.3), Inches(1), align=PP_ALIGN.CENTER)
    
    rect(slide, Inches(3.5), Inches(1.5), Inches(6.3), Inches(4.5), None, line=ACCENT_BLUE)
    txt(slide, "Application", Inches(3.5), Inches(1.6), Inches(6.3), Inches(0.5), align=PP_ALIGN.CENTER)
    
    rect(slide, Inches(2), Inches(0.5), Inches(9.3), Inches(6.5), None, line=TEXT_MUTED)
    txt(slide, "Infrastructure (adapters, config)", Inches(2), Inches(0.6), Inches(9.3), Inches(0.5), align=PP_ALIGN.CENTER)

    footer(slide, "B3")
    return slide

def slide_b4_composition(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Backup: Capability composition")
    
    card(slide, Inches(2), Inches(2), Inches(4), Inches(4), "Always Included", ["core", "api", "security", "logging", "observability", "mapping", "test"], ACCENT_BLUE)
    card(slide, Inches(7), Inches(2), Inches(4), Inches(4), "Optional (On-Demand)", ["data", "nats", "http-client", "soap-client"], ACCENT_GREEN)
    
    footer(slide, "B4")
    return slide

def slide_b5_security(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Backup: Security model")
    
    card(slide, Inches(1), Inches(3), Inches(4), Inches(2.5), "Gravitee API Gateway", ["Gateway routing", "Policies & rate limiting", "Access control"], TEXT_MUTED)
    txt(slide, "--- JWT Token --->", Inches(5), Inches(4), Inches(3), Inches(0.5), size=14, color=ACCENT_BLUE, align=PP_ALIGN.CENTER)
    card(slide, Inches(8), Inches(3), Inches(4), Inches(2.5), "Spring Boot Service", ["OAuth2 Resource Server", "JWT Validation", "Roles / scopes"], ACCENT_GREEN)
    
    footer(slide, "B5")
    return slide

def slide_b6_deployment(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Backup: Deployment boundary")
    
    card(slide, Inches(1), Inches(2), Inches(5), Inches(3.5), "Inside foundation-platform", ["parent", "BOM", "starters", "common utilities", "archetype templates"], ACCENT_BLUE)
    card(slide, Inches(7), Inches(2), Inches(5), Inches(3.5), "Outside foundation-platform", ["Helm charts", "Kubernetes manifests", "Runtime CI/CD", "Service repositories"], TEXT_MUTED)
    txt(slide, "The foundation is a development accelerator, not a deployment platform.", Inches(1), Inches(6), Inches(11), Inches(0.5), size=16, color=ACCENT_GREEN, italic=True)
    footer(slide, "B6")
    return slide

def slide_b7_integration(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Backup: Integration model")
    
    multiline(slide, ["REST: OpenAPI contract → generated client → WebClient → adapter",
                      "SOAP: WSDL contract → generated CXF client → SOAP runtime → adapter",
                      "Messaging: NATS subject → envelope → consumer → idempotence"], Inches(1), Inches(2.5), Inches(11), Inches(3), size=16, line_gap=16)
    footer(slide, "B7")
    return slide

def slide_b8_risks(prs):
    slide = blank(prs)
    bg(slide, BG_DEEP_NAVY)
    header(slide, "Backup: Risks & Mitigations")
    
    lines = [
        "1. Heavy framework perception → Kept Spring Boot-native with clear overrides",
        "2. Teams bypass the foundation → Archetype DX made faster than manual setup",
        "3. Creeping mandatory features → Maintained strict dependency-driven composition",
        "4. Leaking generated clients → Mapped via infrastructure adapters (hexagonal)"
    ]
    multiline(slide, lines, Inches(1), Inches(2.5), Inches(11), Inches(3), size=14, line_gap=16)
    footer(slide, "B8")
    return slide

# ─────────────────────────────────────────────
# MAIN EXECUTION
# ─────────────────────────────────────────────
def main():
    prs = new_prs()
    
    # Main Deck (1-15)
    slide_1_title(prs)
    slide_2_challenge(prs)
    slide_3_cost(prs)
    slide_4_foundation(prs)
    slide_5_principles(prs)
    slide_6_architecture(prs)
    slide_7_dx(prs)
    slide_8_mandatory(prs)
    slide_9_optional(prs)
    slide_10_migration(prs)
    slide_11_governance(prs)
    slide_12_productivity(prs)
    slide_13_boundaries(prs)
    slide_14_adoption(prs)
    slide_15_closing(prs)
    
    # Backup Deck (B1-B8)
    slide_b1_map(prs)
    slide_b2_archetype(prs)
    slide_b3_hexagonal(prs)
    slide_b4_composition(prs)
    slide_b5_security(prs)
    slide_b6_deployment(prs)
    slide_b7_integration(prs)
    slide_b8_risks(prs)

    output_path = "foundation_platform_presentation_consulting_deck.pptx"
    prs.save(output_path)
    print(f"✅ Deck généré avec succès dans : {output_path}")

if __name__ == "__main__":
    main()
